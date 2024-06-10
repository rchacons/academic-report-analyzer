from ..schemas.comparison_schema import Subject, Book
from typing import List
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
import xlsxwriter


class ExportService:
    
    def create_excel_subjects(self, subjects: List[Subject], file_name: str):
        """
        This method creates an Excel file with a table that contains information about subjects.
        
        Parameters:
        subjects (List[Subject]): A list of Subject objects that contain the data to be written to the Excel file.
        file_name (str): The name of the Excel file to be created.
        """
        data = []
        for subject in subjects:
            for config in subject.materials_configurations:
                materials = ', '.join(config.materials)
                row = {
                    'Domaine': subject.field,
                    'Niveau': subject.level,
                    'Intitulé': subject.title,
                    'Matériels': materials,
                    'Thème': subject.theme
                }
                data.append(row)
        self._create_excel_file(data, file_name)  

    def create_excel_books(self, books: List[Book], file_name: str):
        """
        This method creates an Excel file with a table that contains information about books.
        
        Parameters:
        books (List[Books]): A list of Books objects that contain the data to be written to the Excel file.
        file_name (str): The name of the Excel file to be created.
        """
        data = []
        for book in books:
            row = {
                    'Titre': book.book_name,
                    'Auteur': book.author,
                    'Date de Publication': book.year_published
                }
            data.append(row)
            
        self._create_excel_file(data, file_name)        
        
    def _create_excel_file(self, data: List[dict], file_name: str):
        """
        This method creates an Excel file with a table that contains information about subjects or books.
        
        Parameters:
        data (List[dict]): A list of dictionaries that contain the data to be written to the Excel file.
        file_name (str): The name of the Excel file to be created.
        """
        df = pd.DataFrame(data)
       
        # Create a new Excel file and add a worksheet
        workbook = xlsxwriter.Workbook(file_name)
        worksheet = workbook.add_worksheet()

        # Add a table to the worksheet
        worksheet.add_table(0, 0, len(df), len(df.columns) - 1, {
            'data': df.values.tolist(),
            'columns': [{'header': column} for column in df.columns],
            'style': 'Table Style Medium 9',
            'autofilter': True,
        })

        # Set the column width and format
        for i, column in enumerate(df.columns):
            worksheet.set_column(i, i, max(df[column].astype(str).apply(len)) + 2)

        workbook.close()